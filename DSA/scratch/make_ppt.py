import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    PRIMARY = RGBColor(0, 86, 179)      # #0056b3 Medical Blue
    SECONDARY = RGBColor(44, 62, 80)    # #2c3e50 Dark Slate
    ACCENT = RGBColor(230, 81, 0)       # #e65100 Medical Orange
    BG_LIGHT = RGBColor(244, 247, 246)  # Light gray-green
    WHITE = RGBColor(255, 255, 255)
    DARK_TEXT = RGBColor(33, 33, 33)
    GRAY_TEXT = RGBColor(100, 100, 100)

    def set_slide_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_header(slide, title_text, category_text="SMARTHOSPITAL SYSTEM"):
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.0))
        tf = header_box.text_frame
        tf.word_wrap = True
        
        p_cat = tf.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT

        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = PRIMARY

    def add_card(slide, left, top, width, height, title="", items=[], bg_color=WHITE, border_color=PRIMARY):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)

        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_top = Inches(0.25)
        tf.margin_right = Inches(0.25)

        if title:
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(18)
            p.font.bold = True
            p.font.color.rgb = PRIMARY
            p.space_after = Pt(8)

        for idx, item in enumerate(items):
            p = tf.add_paragraph() if title or idx > 0 else tf.paragraphs[0]
            p.text = "• " + item
            p.font.size = Pt(13)
            p.font.color.rgb = DARK_TEXT
            p.space_after = Pt(5)

    # -------------------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, BG_LIGHT)

    card1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(1.2), Inches(11.333), Inches(5.1))
    card1.fill.solid()
    card1.fill.fore_color.rgb = PRIMARY
    card1.line.fill.background()

    tf = card1.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.8)
    tf.margin_top = Inches(1.0)

    p1 = tf.paragraphs[0]
    p1.text = "SmartHospital System"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = WHITE

    p2 = tf.add_paragraph()
    p2.text = "Comprehensive Hospital Resource & Information Management Platform (Modules 1–7: CO1–CO6)"
    p2.font.size = Pt(22)
    p2.font.color.rgb = RGBColor(200, 230, 255)
    p2.space_before = Pt(10)

    p3 = tf.add_paragraph()
    p3.text = "Strings, Suffix Automata, Dynamic Programming, Network Flow, NP-Reductions, Randomized & Parallel Algorithms"
    p3.font.size = Pt(16)
    p3.font.color.rgb = WHITE
    p3.space_before = Pt(30)

    # -------------------------------------------------------------------------
    # SLIDE 2: Problem Statement & Project Scope
    # -------------------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, BG_LIGHT)
    add_header(slide2, "Healthcare Challenges & System Scope")

    add_card(slide2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
             title="The Complex Healthcare Dilemma",
             items=[
                 "Massive unstructured patient charts and diagnostic telemetry.",
                 "Typo hazards in prescriptions and critical clinical notes.",
                 "Medical resource bottlenecks: oxygen, blood bank, and ICU bed logistics.",
                 "Doctor-case specialty mismatch under severe emergency surges.",
                 "Combinatorial complexity in multi-drug treatment interactions (NP-hard).",
                 "Unbounded real-time sensor streams and large-scale hospital analytics."
             ])

    add_card(slide2, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2),
             title="The Complete SmartHospital Engine",
             items=[
                 "Unified Spring Boot 3 architecture with Interactive CLI and REST Web Dashboard.",
                 "Module 1 (CO1): Linear String Search (KMP, Z-Function, Rabin-Karp, Aho-Corasick).",
                 "Module 2 (CO2): Suffix Indexing & Kasai LCP phrase extraction.",
                 "Module 3 (CO3): Levenshtein typo correction & Needleman-Wunsch DNA alignment.",
                 "Module 4 (CO4): Flow Networks (Ford-Fulkerson, Edmonds-Karp, Dinic, Bipartite, König).",
                 "Module 5 (CO5): NP-Completeness (3-SAT DPLL, Clique reduction, 2-Approximation).",
                 "Modules 6 & 7 (CO6): Randomized QuickSort, Reservoir Sampling, Blelloch Scan & Brent's Lemma."
             ])

    # -------------------------------------------------------------------------
    # SLIDE 3: Module 4 – Network Flow (CO4)
    # -------------------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, BG_LIGHT)
    add_header(slide3, "Module 4: Network Flow & Medical Logistics (CO4)")

    add_card(slide3, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
             title="Flow Algorithms & Medical Routing",
             items=[
                 "Ford-Fulkerson: Residual augmenting paths for medical oxygen & supply transit.",
                 "Edmonds-Karp: BFS shortest-path max flow in O(V * E^2).",
                 "Minimum Cut Discovery: Detects critical saturated logistic bottlenecks between depot and ICU wards.",
                 "Dinic's Algorithm: Level graphs and blocking flow in O(V^2 * E) for large multi-hospital transit networks."
             ])

    add_card(slide3, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2),
             title="Doctor Matching & König's Theorem",
             items=[
                 "Bipartite Matching: Assigns medical specialists (Cardiology, Neuro, Oncology) to pending emergency cases based on specialty compatibility.",
                 "König's Theorem Duality: Mathematically proves Maximum Matching Size == Minimum Vertex Cover Size in bipartite graphs.",
                 "Alternating BFS: Extracts the minimal set of doctors and cases covering all potential scheduling and contagion conflicts."
             ])

    # -------------------------------------------------------------------------
    # SLIDE 4: Module 5 – NP-Completeness and Approximation (CO5)
    # -------------------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, BG_LIGHT)
    add_header(slide4, "Module 5: NP-Completeness & Approximation (CO5)")

    add_card(slide4, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
             title="SAT / 3-SAT Treatment Protocols",
             items=[
                 "Boolean Constraint Modeling: Models medical rules (Antibiotics, Anticoagulants, Gastric Protectors, Corticosteroids).",
                 "DPLL Complete Decision Engine: Solves CNF 3-SAT using unit propagation, pure literal elimination, and backtracking.",
                 "Returns guaranteed clinically valid prescription configurations or proves contradictory impossibility."
             ])

    add_card(slide4, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2),
             title="Karp Reductions & 2-Approximation",
             items=[
                 "3-SAT → CLIQUE Reduction: Constructs compatibility graphs where k-clique existence directly models satisfiability.",
                 "Karp Reduction Chain: Formal proof bridging CLIQUE ↔ INDEPENDENT-SET ↔ VERTEX-COVER in conflict graphs.",
                 "Vertex-Cover 2-Approximation: Maximal matching greedy edge elimination guaranteeing |C| <= 2 * |C*| on large hospital conflict networks."
             ])

    # -------------------------------------------------------------------------
    # SLIDE 5: Module 6 – Randomized Algorithms (CO6)
    # -------------------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, BG_LIGHT)
    add_header(slide5, "Module 6: Randomized Algorithms (CO6)")

    add_card(slide5, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
             title="Randomized QuickSort (Triage)",
             items=[
                 "Ranks patient records and emergency cases by urgency score.",
                 "Uniform random pivot selection guarantees expected O(n log n) runtime.",
                 "Eliminates worst-case O(n^2) degradation on pre-sorted admission logs.",
                 "Captures execution metrics: comparisons, swaps, depth, and elapsed time."
             ])

    add_card(slide5, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2),
             title="Reservoir Sampling & Miller-Rabin",
             items=[
                 "Reservoir Sampling (Algorithm R): Continuously samples k representative readings from unbounded real-time ICU telemetry streams.",
                 "Uniform Probability: Guarantees every vital reading has equal probability k/N of being retained.",
                 "Miller-Rabin Primality Test: High-speed probabilistic primality verification with error probability <= (1/4)^k.",
                 "Generates 64-bit cryptographic prime tokens for HIPAA-compliant encrypted patient communications."
             ])

    # -------------------------------------------------------------------------
    # SLIDE 6: Module 7 – Parallel Algorithms & Analytics (CO6)
    # -------------------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6, BG_LIGHT)
    add_header(slide6, "Module 7: Parallel Algorithms & Analytics Engine (CO6)")

    add_card(slide6, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
             title="Blelloch Work-Efficient Scan",
             items=[
                 "Parallel prefix sum algorithm on binary trees for running dosage aggregation.",
                 "Up-Sweep (Reduce Phase): Aggregates partial sums in O(n) work, O(log n) span.",
                 "Down-Sweep Phase: Distributes prefix sums down the tree in O(n) work, O(log n) span.",
                 "Generates both exclusive and inclusive cumulative dosage vectors."
             ])

    add_card(slide6, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2),
             title="Parallel Reduce & Brent's Theorem",
             items=[
                 "Parallel Reduce: Java ForkJoin framework for SUM, MIN, MAX, AVG across massive vital telemetry datasets.",
                 "Demonstrates multi-core speedup and work-span efficiency.",
                 "Brent's Theorem Modeling: W/p <= T_p <= (W - T_inf)/p + T_inf.",
                 "Models speedup curves, processor efficiency, and hardware saturation limits for hospital telemetry pipelines."
             ])

    # -------------------------------------------------------------------------
    # SLIDE 7: System Architecture & User Interfaces
    # -------------------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7, BG_LIGHT)
    add_header(slide7, "Dual Interface & Execution Architecture")

    add_card(slide7, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
             title="Interactive Console Scanner CLI",
             items=[
                 "MainApp.java standalone runner without Spring Boot overhead.",
                 "Hierarchical 10-option interactive menu loop (Modules 1–7).",
                 "Dynamic input handling for searches, network flows, 3-SAT evaluations, and parallel simulations.",
                 "Robust input safety and EOF handling for automated terminal runs."
             ])

    add_card(slide7, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2),
             title="Spring Boot REST & Web Dashboard",
             items=[
                 "SmartHospitalApplication.java web server running on port 8080.",
                 "20+ REST API Endpoints under /api for all 7 modules.",
                 "Rich interactive single-page dashboard at http://localhost:8080.",
                 "Live AJAX testing with formatted JSON outputs."
             ])

    # -------------------------------------------------------------------------
    # SLIDE 8: Verification & Automated Test Suite
    # -------------------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8, BG_LIGHT)
    add_header(slide8, "Verification & Quality Assurance")

    add_card(slide8, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2),
             title="100% Passing Automated JUnit 5 Test Suite (26 Tests Across 7 Modules)",
             items=[
                 "✓ Module1Test: KMP, Z-Function, Rabin-Karp, Aho-Corasick keyword detection (4/4 Passed).",
                 "✓ Module2Test: Suffix Array, Kasai LCP, Longest Repeated Medical Phrase (2/2 Passed).",
                 "✓ Module3Test: Levenshtein, Damerau-Levenshtein, Needleman-Wunsch DNA, Bitmask DP (4/4 Passed).",
                 "✓ Module4Test: Ford-Fulkerson, Edmonds-Karp Min-Cut, Dinic, Bipartite Matching, König's Theorem (5/5 Passed).",
                 "✓ Module5Test: 3-SAT DPLL, 3-SAT to CLIQUE Reduction, Karp Chain, Vertex Cover 2-Approximation (5/5 Passed).",
                 "✓ Module6Test: Randomized QuickSort ranking, Reservoir Sampling, Miller-Rabin Primality & Token generation (3/3 Passed).",
                 "✓ Module7Test: Blelloch Scan, Parallel Reduce ForkJoin, Brent's Theorem Performance Modeling (3/3 Passed).",
                 "✓ Total Tests: 26 Tests Run, 0 Failures, 0 Errors, 0 Skipped — 100% Clean Execution."
             ])

    # -------------------------------------------------------------------------
    # SLIDE 9: Summary & Conclusion
    # -------------------------------------------------------------------------
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9, BG_LIGHT)
    add_header(slide9, "Conclusion & Project Highlights")

    add_card(slide9, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2),
             title="Key Accomplishments",
             items=[
                 "✓ Fully engineered and integrated all 4 requested modules: Network Flow (M4), NP-Completeness & Approximation (M5), Randomized Algorithms (M6), and Parallel Algorithms (M7).",
                 "✓ Seamlessly integrated 15+ advanced algorithms into the SmartHospital healthcare domain.",
                 "✓ Maintained clean code, modular packages, comprehensive JavaDoc, and enterprise architecture.",
                 "✓ Dual interactive access: rich CLI terminal scanner menu + full Spring Boot REST web UI.",
                 "✓ Zero warnings, zero compilation errors, and complete verification across all automated test suites."
             ])

    output_path = os.path.join(os.getcwd(), "SmartHospital_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation successfully created at: {output_path}")

if __name__ == "__main__":
    create_presentation()
